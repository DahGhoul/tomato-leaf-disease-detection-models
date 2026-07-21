import io
import json
from datetime import datetime

DISEASE_NAMES = {
    'Tomato___Bacterial_spot': 'Mancha Bacteriana',
    'Tomato___Early_blight': 'Tizón Temprano',
    'Tomato___Late_blight': 'Tizón Tardío',
    'Tomato___Leaf_Mold': 'Moho de Hoja',
    'Tomato___Septoria_leaf_spot': 'Mancha de Septoria',
    'Tomato___Spider_mites Two-spotted_spider_mite': 'Ácaros Araña',
    'Tomato___Target_Spot': 'Mancha Diana',
    'Tomato___Tomato_Yellow_Leaf_Curl_Virus': 'Virus del Rizado Amarillo',
    'Tomato___Tomato_mosaic_virus': 'Virus del Mosaico',
    'Tomato___healthy': 'Saludable',
}

# Keys that are NOT model predictions in the result dict
NON_MODEL_KEYS = {"ensemble_prediction", "mean_confidence"}


def _get_model_entries(predictions: dict):
    """Extract only model result entries, ignoring ensemble/meta keys."""
    return {k: v for k, v in predictions.items() if k not in NON_MODEL_KEYS and isinstance(v, dict)}


def _friendly_disease(raw: str) -> str:
    return DISEASE_NAMES.get(raw, raw.replace('Tomato___', '').replace('_', ' '))


def generate_pdf_report(predictions: dict, image_buffer=None):
    """Generate a comprehensive PDF diagnostic report."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=0.6*inch, bottomMargin=0.6*inch,
                            leftMargin=0.8*inch, rightMargin=0.8*inch)
    styles = getSampleStyleSheet()
    story = []

    # ── Title ──────────────────────────────────────────────────────────
    title_style = ParagraphStyle('Title', parent=styles['Title'], fontSize=20,
                                 textColor=colors.HexColor('#1a202c'), spaceAfter=6, alignment=TA_CENTER)
    sub_style = ParagraphStyle('Sub', parent=styles['Normal'], fontSize=10,
                               textColor=colors.HexColor('#718096'), alignment=TA_CENTER)
    story.append(Paragraph("🍅 Reporte de Diagnóstico — Detección de Enfermedades en Hojas de Tomate", title_style))
    story.append(Paragraph(f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}", sub_style))
    story.append(Spacer(1, 0.3*inch))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#e2e8f0')))
    story.append(Spacer(1, 0.2*inch))

    # ── Uploaded image ─────────────────────────────────────────────────
    if image_buffer:
        try:
            image_buffer.seek(0)
            img = RLImage(image_buffer, width=2.5*inch, height=2.5*inch)
            img_table = Table([[img]], colWidths=[2.5*inch])
            img_table.setStyle(TableStyle([('ALIGN', (0, 0), (-1, -1), 'CENTER')]))
            story.append(img_table)
            story.append(Spacer(1, 0.2*inch))
        except Exception:
            pass

    # ── Ensemble consensus ─────────────────────────────────────────────
    ensemble = predictions.get("ensemble_prediction", "N/D")
    mean_conf = predictions.get("mean_confidence", 0)
    consensus_name = _friendly_disease(ensemble)

    h2 = ParagraphStyle('H2', parent=styles['Heading2'], fontSize=13, textColor=colors.HexColor('#2d3748'))
    story.append(Paragraph("Diagnóstico por Consenso de Modelos", h2))
    consensus_data = [
        ["Enfermedad detectada", "Confianza promedio"],
        [consensus_name, f"{mean_conf * 100:.2f}%"],
    ]
    ct = Table(consensus_data, colWidths=[3.5*inch, 2*inch])
    ct.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#38a169')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTSIZE', (0, 0), (-1, -1), 11),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#f0fff4')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#c6f6d5')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f0fff4')]),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(ct)
    story.append(Spacer(1, 0.3*inch))

    # ── Per-model table ────────────────────────────────────────────────
    story.append(Paragraph("Resultados por Modelo", h2))
    model_entries = _get_model_entries(predictions)
    table_data = [["Modelo", "Diagnóstico", "Confianza", "Tiempo (s)"]]
    for model_name, res in model_entries.items():
        table_data.append([
            model_name,
            _friendly_disease(res.get("prediction", "")),
            f"{res.get('confidence', 0) * 100:.2f}%",
            f"{res.get('inference_time', 0):.3f}s",
        ])

    mt = Table(table_data, colWidths=[1.8*inch, 2.2*inch, 1.2*inch, 1.0*inch])
    mt.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2b6cb0')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#bee3f8')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#ebf8ff')]),
        ('TOPPADDING', (0, 0), (-1, -1), 7),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
    ]))
    story.append(mt)
    story.append(Spacer(1, 0.4*inch))

    # ── Statistical notes ──────────────────────────────────────────────
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#e2e8f0')))
    story.append(Spacer(1, 0.15*inch))
    note_style = ParagraphStyle('Note', parent=styles['Normal'], fontSize=8,
                                textColor=colors.HexColor('#718096'))
    story.append(Paragraph(
        "Este reporte fue generado automáticamente por el sistema de detección de enfermedades de hojas de tomate "
        "usando un ensamble de modelos de aprendizaje profundo (MobileNetV3, EfficientNetB7, ResNet50) "
        "y modelos híbridos (SVM + MobileNet, Random Forest + EfficientNet).", note_style))

    doc.build(story)
    buffer.seek(0)
    return buffer


def generate_word_report(predictions: dict, image_buffer=None):
    """Generate a Word (.docx) diagnostic report."""
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Title
    title = doc.add_heading('🍅 Reporte de Diagnóstico - Detección de Enfermedades en Hojas de Tomate', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f'Fecha de generación: {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}')
    doc.add_paragraph('')

    # Image
    if image_buffer:
        try:
            image_buffer.seek(0)
            doc.add_picture(image_buffer, width=Inches(3))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            pass

    # Ensemble
    doc.add_heading('Diagnóstico por Consenso', level=1)
    ensemble = predictions.get("ensemble_prediction", "N/D")
    mean_conf = predictions.get("mean_confidence", 0)
    p = doc.add_paragraph()
    p.add_run('Enfermedad: ').bold = True
    p.add_run(_friendly_disease(ensemble))
    p.add_run('\nConfianza promedio: ').bold = True
    p.add_run(f'{mean_conf * 100:.2f}%')

    # Per model
    doc.add_heading('Resultados por Modelo', level=1)
    model_entries = _get_model_entries(predictions)

    table = doc.add_table(rows=1, cols=4)
    table.style = 'Light Shading Accent 1'
    hdr = table.rows[0].cells
    hdr[0].text = 'Modelo'
    hdr[1].text = 'Diagnóstico'
    hdr[2].text = 'Confianza'
    hdr[3].text = 'Tiempo (s)'

    for model_name, res in model_entries.items():
        row = table.add_row().cells
        row[0].text = model_name
        row[1].text = _friendly_disease(res.get('prediction', ''))
        row[2].text = f"{res.get('confidence', 0) * 100:.2f}%"
        row[3].text = f"{res.get('inference_time', 0):.3f}s"

    doc.add_paragraph('')
    doc.add_paragraph(
        'Este reporte fue generado automáticamente por el sistema de detección de enfermedades de hojas de tomate.',
        style='Caption'
    )

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def generate_excel_report(predictions: dict):
    """Generate an Excel (.xlsx) diagnostic report."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Diagnóstico"

    # ── Styling helpers ──────────────────────────────────────────────
    header_fill = PatternFill("solid", fgColor="2B6CB0")
    header_font = Font(color="FFFFFF", bold=True)
    accent_fill = PatternFill("solid", fgColor="EBF8FF")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style='thin', color='BEE3F8')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ── Title ─────────────────────────────────────────────────────────
    ws.merge_cells('A1:F1')
    ws['A1'] = '🍅 Reporte de Diagnóstico — Detección de Enfermedades en Hojas de Tomate'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = center
    ws.row_dimensions[1].height = 30

    ws.merge_cells('A2:F2')
    ws['A2'] = f'Generado: {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}'
    ws['A2'].alignment = center
    ws['A2'].font = Font(italic=True, color='718096')
    ws.row_dimensions[2].height = 18

    # ── Ensemble row ──────────────────────────────────────────────────
    ws['A4'] = 'DIAGNÓSTICO POR CONSENSO'
    ws['A4'].font = Font(bold=True, size=12, color='276749')
    ws['B4'] = _friendly_disease(predictions.get("ensemble_prediction", "N/D"))
    ws['C4'] = 'CONFIANZA PROMEDIO'
    ws['C4'].font = Font(bold=True)
    ws['D4'] = f"{predictions.get('mean_confidence', 0) * 100:.2f}%"

    # ── Per-model table ────────────────────────────────────────────────
    headers = ['Modelo', 'Diagnóstico (EN)', 'Diagnóstico (ES)', 'Confianza (%)', 'Tiempo Inf. (s)', 'Entropía']
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=6, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[6].height = 22

    model_entries = _get_model_entries(predictions)
    for row_idx, (model_name, res) in enumerate(model_entries.items(), start=7):
        fill = PatternFill("solid", fgColor="EBF8FF") if row_idx % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
        raw_pred = res.get('prediction', '')
        values = [
            model_name,
            raw_pred,
            _friendly_disease(raw_pred),
            round(res.get('confidence', 0) * 100, 2),
            round(res.get('inference_time', 0), 4),
            round(res.get('entropy', 0), 4),
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.fill = fill
            cell.alignment = center
            cell.border = border

    # ── Column widths ─────────────────────────────────────────────────
    col_widths = [22, 40, 28, 18, 18, 14]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
